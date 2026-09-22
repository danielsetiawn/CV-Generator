from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_TEMPLATE = PROJECT_ROOT / "templates" / "Accessible-MCS-Resume-Template-Bullet-Points.docx"

ACTION_VERBS = {
    "accelerated",
    "accomplished",
    "achieved",
    "acquired",
    "administered",
    "advised",
    "aligned",
    "allocated",
    "analyzed",
    "approved",
    "architected",
    "assessed",
    "assigned",
    "assisted",
    "attained",
    "audited",
    "authored",
    "automated",
    "benchmarked",
    "boosted",
    "briefed",
    "built",
    "calculated",
    "chaired",
    "coached",
    "coded",
    "collaborated",
    "communicated",
    "compiled",
    "composed",
    "conducted",
    "configured",
    "contributed",
    "conveyed",
    "coordinated",
    "corresponded",
    "created",
    "debugged",
    "decreased",
    "delegated",
    "delivered",
    "deployed",
    "designed",
    "developed",
    "directed",
    "discovered",
    "doubled",
    "drafted",
    "drove",
    "earned",
    "edited",
    "eliminated",
    "enabled",
    "engineered",
    "enhanced",
    "ensured",
    "established",
    "evaluated",
    "examined",
    "exceeded",
    "executed",
    "expanded",
    "expedited",
    "explored",
    "fabricated",
    "facilitated",
    "forecasted",
    "formulated",
    "founded",
    "generated",
    "guided",
    "handled",
    "headed",
    "hired",
    "identified",
    "illustrated",
    "implemented",
    "improved",
    "increased",
    "influenced",
    "informed",
    "initiated",
    "inspected",
    "installed",
    "instructed",
    "integrated",
    "interpreted",
    "introduced",
    "investigated",
    "launched",
    "led",
    "lowered",
    "maintained",
    "managed",
    "marketed",
    "maximized",
    "measured",
    "mediated",
    "mentored",
    "migrated",
    "minimized",
    "modeled",
    "modified",
    "monitored",
    "motivated",
    "navigated",
    "negotiated",
    "obtained",
    "optimized",
    "organized",
    "originated",
    "oversaw",
    "partnered",
    "performed",
    "persuaded",
    "piloted",
    "pioneered",
    "planned",
    "prepared",
    "presented",
    "prioritized",
    "produced",
    "programmed",
    "promoted",
    "publicized",
    "published",
    "quantified",
    "recruited",
    "reduced",
    "refactored",
    "regulated",
    "reorganized",
    "repaired",
    "represented",
    "researched",
    "resolved",
    "restructured",
    "revamped",
    "saved",
    "secured",
    "served",
    "set",
    "spearheaded",
    "standardized",
    "streamlined",
    "strengthened",
    "supervised",
    "supported",
    "surpassed",
    "surveyed",
    "tested",
    "trained",
    "transformed",
    "translated",
    "troubleshot",
    "upgraded",
    "wrote",
    "yielded",
}

PRONOUN_RE = re.compile(r"\b(i|me|my|mine|we|our|ours)\b", re.IGNORECASE)


class CVGeneratorError(ValueError):
    """Raised when resume data is incomplete or invalid."""


def load_cv_data(path: str | Path) -> dict[str, Any]:
    """Load resume data from JSON, or YAML when PyYAML is installed."""
    source = Path(path)
    if not source.exists():
        raise CVGeneratorError(f"Data file not found: {source}")

    text = source.read_text(encoding="utf-8")
    suffix = source.suffix.lower()
    if suffix == ".json":
        return json.loads(text)
    if suffix in {".yaml", ".yml"}:
        try:
            import yaml  # type: ignore
        except ImportError as exc:
            raise CVGeneratorError("Install PyYAML to read YAML files, or use JSON input.") from exc
        loaded = yaml.safe_load(text)
        return loaded or {}
    raise CVGeneratorError("Use a .json, .yaml, or .yml data file.")


def generate_cv(
    data: dict[str, Any],
    output_path: str | Path,
    template_path: str | Path = DEFAULT_TEMPLATE,
) -> Path:
    """Generate an ATS-friendly Harvard-style resume as a DOCX file."""
    template = Path(template_path)
    if not template.exists():
        raise CVGeneratorError(f"Template file not found: {template}")

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)

    doc = Document(str(template))
    _clear_body(doc)
    _configure_document(doc)

    profile = _require_mapping(data, "profile")
    _add_profile(doc, profile)

    for section in data.get("section_order", _default_section_order(data)):
        key = str(section).lower()
        if key == "education":
            _add_education(doc, data.get("education", []))
        elif key == "experience":
            _add_experience_section(doc, "Experience", data.get("experience", []))
        elif key == "leadership":
            _add_experience_section(doc, "Leadership & Activities", data.get("leadership", []))
        elif key == "projects":
            _add_experience_section(doc, "Projects", data.get("projects", []))
        elif key == "skills":
            _add_skills(doc, data.get("skills", {}), data.get("interests"))

    doc.save(str(output))
    return output


def ats_warnings(data: dict[str, Any]) -> list[str]:
    """Return practical warnings for common ATS resume issues."""
    warnings: list[str] = []
    profile = data.get("profile", {})
    if not profile.get("name"):
        warnings.append("Profile name is missing.")
    if not profile.get("email"):
        warnings.append("Email is missing.")
    if not profile.get("phone"):
        warnings.append("Phone number is missing.")

    for section_key in ("experience", "leadership", "projects"):
        for idx, item in enumerate(_as_list(data.get(section_key, [])), start=1):
            for bullet in _as_list(item.get("bullets", [])):
                cleaned = str(bullet).strip()
                if not cleaned:
                    continue
                first = re.sub(r"[^A-Za-z]", "", cleaned.split()[0]).lower()
                if first and first not in ACTION_VERBS:
                    warnings.append(
                        f"{section_key} item {idx} bullet may not start with a strong action verb: {cleaned[:70]}"
                    )
                if PRONOUN_RE.search(cleaned):
                    warnings.append(f"{section_key} item {idx} bullet uses a personal pronoun: {cleaned[:70]}")
                if len(cleaned) > 220:
                    warnings.append(f"{section_key} item {idx} bullet is long; consider trimming below 220 characters.")
    return warnings


def _clear_body(doc: Document) -> None:
    body = doc._body._element
    for child in list(body):
        if child.tag.endswith("sectPr"):
            continue
        body.remove(child)


def _configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.top_margin = Inches(0.5)
    section.right_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.5)

    _set_style(doc, "Normal", size=10.5)
    _set_style(doc, "Heading 1", size=16, bold=True)
    _set_style(doc, "Heading 2", size=11, bold=True)
    _set_style(doc, "List Bullet", size=10.5)
    _set_style(doc, "List Paragraph", size=10.5)

    normal = doc.styles["Normal"].paragraph_format
    normal.alignment = WD_ALIGN_PARAGRAPH.LEFT
    normal.left_indent = Inches(0)
    normal.right_indent = Inches(0)
    normal.first_line_indent = Inches(0)
    normal.space_before = Pt(0)
    normal.space_after = Pt(0)
    normal.line_spacing = 1.0
    normal.line_spacing_rule = WD_LINE_SPACING.SINGLE

    heading_1 = doc.styles["Heading 1"].paragraph_format
    heading_1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading_1.left_indent = Inches(0)
    heading_1.right_indent = Inches(0)
    heading_1.first_line_indent = Inches(0)
    heading_1.space_before = Pt(0)
    heading_1.space_after = Pt(2)
    heading_1.line_spacing_rule = WD_LINE_SPACING.SINGLE

    heading_2 = doc.styles["Heading 2"].paragraph_format
    heading_2.alignment = WD_ALIGN_PARAGRAPH.LEFT
    heading_2.left_indent = Inches(0)
    heading_2.right_indent = Inches(0)
    heading_2.first_line_indent = Inches(0)
    heading_2.space_before = Pt(7)
    heading_2.space_after = Pt(2)
    heading_2.line_spacing_rule = WD_LINE_SPACING.SINGLE

    for style_name in ("List Bullet", "List Paragraph"):
        if style_name in doc.styles:
            fmt = doc.styles[style_name].paragraph_format
            fmt.alignment = WD_ALIGN_PARAGRAPH.LEFT
            fmt.space_before = Pt(0)
            fmt.space_after = Pt(0)
            fmt.left_indent = Inches(0.25)
            fmt.right_indent = Inches(0)
            fmt.line_spacing_rule = WD_LINE_SPACING.SINGLE


def _set_style(doc: Document, name: str, size: float, bold: bool | None = None) -> None:
    try:
        style = doc.styles[name]
    except KeyError:
        style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold


def _add_profile(doc: Document, profile: dict[str, Any]) -> None:
    name = str(profile.get("name") or "Your Name").strip()
    name_para = doc.add_paragraph(style="Heading 1")
    _set_paragraph_layout(name_para, alignment=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=2)
    name_run = name_para.add_run(name)
    name_run.bold = True

    contact_parts = [
        profile.get("location"),
        profile.get("email"),
        profile.get("phone"),
        profile.get("linkedin"),
        profile.get("website"),
    ]
    contact = " | ".join(str(part).strip() for part in contact_parts if part)
    if contact:
        p = doc.add_paragraph(style="Normal")
        _set_paragraph_layout(p, alignment=WD_ALIGN_PARAGRAPH.CENTER)
        p.add_run(contact)


def _add_education(doc: Document, items: Any) -> None:
    entries = _as_list(items)
    if not entries:
        return
    _section_heading(doc, "Education")
    for item in entries:
        school = str(item.get("school") or "").strip()
        location = str(item.get("location") or "").strip()
        degree = str(item.get("degree") or "").strip()
        date = str(item.get("graduation_date") or item.get("dates") or "").strip()
        detail_parts = [item.get("concentration"), _prefix("GPA", item.get("gpa"))]

        if not school and not degree and not item.get("bullets"):
            continue

        if school or location:
            _tabbed_line(doc, school or degree, location, bold_left=True, bold_right=True)
            if school and (degree or any(detail_parts) or date):
                _tabbed_line(doc, _join_sentence_parts([degree, *detail_parts]), date, italic_left=True)
        elif degree or date or any(detail_parts):
            _tabbed_line(doc, _join_sentence_parts([degree, *detail_parts]), date, italic_left=True)

        for label in ("thesis", "relevant_coursework", "honors"):
            value = item.get(label)
            if value:
                _label_line(doc, _title_label(label), _stringify_list(value))
        _bullets(doc, item.get("bullets", []))


def _add_experience_section(doc: Document, title: str, items: Any) -> None:
    entries = _as_list(items)
    if not entries:
        return
    _section_heading(doc, title)
    for item in entries:
        organization = str(item.get("organization") or "").strip()
        location = str(item.get("location") or "").strip()
        role = str(item.get("title") or "").strip()
        dates = str(item.get("dates") or "").strip()
        bullets = item.get("bullets", [])

        if not organization and not role and not bullets:
            continue

        if organization or location:
            _tabbed_line(doc, organization, location, bold_left=True, bold_right=True)
        if role or dates:
            _tabbed_line(doc, role, dates, italic_left=True)
        _bullets(doc, bullets)


def _add_skills(doc: Document, skills: Any, interests: Any = None) -> None:
    if not skills and not interests:
        return
    _section_heading(doc, "Skills & Interests")
    if isinstance(skills, dict):
        for label, values in skills.items():
            _label_line(doc, str(label).replace("_", " ").title(), _stringify_list(values))
    elif skills:
        _label_line(doc, "Skills", _stringify_list(skills))
    if interests:
        _label_line(doc, "Interests", _stringify_list(interests))


def _section_heading(doc: Document, title: str) -> None:
    p = doc.add_paragraph(style="Heading 2")
    _set_paragraph_layout(p, alignment=WD_ALIGN_PARAGRAPH.LEFT, space_before=7, space_after=2)
    _add_bottom_border(p)
    heading_run = p.add_run(title)
    heading_run.bold = True


def _tabbed_line(
    doc: Document,
    left: str,
    right: Any = "",
    *,
    bold_left: bool = False,
    bold_right: bool = False,
    italic_left: bool = False,
    italic_right: bool = False,
) -> None:
    p = doc.add_paragraph(style="Normal")
    _set_paragraph_layout(p, alignment=WD_ALIGN_PARAGRAPH.LEFT)
    p.paragraph_format.tab_stops.add_tab_stop(_usable_width(doc), WD_TAB_ALIGNMENT.RIGHT)
    left_run = p.add_run(str(left).strip())
    left_run.bold = bold_left
    left_run.italic = italic_left
    if right:
        p.add_run("\t")
        right_run = p.add_run(str(right).strip())
        right_run.bold = bold_right
        right_run.italic = italic_right


def _label_line(doc: Document, label: str, value: str) -> None:
    if not value:
        return
    p = doc.add_paragraph(style="Normal")
    _set_paragraph_layout(p, alignment=WD_ALIGN_PARAGRAPH.LEFT)
    lead = p.add_run(f"{label}: ")
    lead.bold = True
    p.add_run(value)


def _bullets(doc: Document, bullets: Any) -> None:
    for bullet in _as_list(bullets):
        text = str(bullet).strip()
        if not text:
            continue
        p = doc.add_paragraph(style="List Bullet")
        _set_paragraph_layout(p, alignment=WD_ALIGN_PARAGRAPH.LEFT)
        p.paragraph_format.left_indent = Inches(0.25)
        p.add_run(text)


def _usable_width(doc: Document) -> int:
    section = doc.sections[0]
    return section.page_width - section.left_margin - section.right_margin


def _default_section_order(data: dict[str, Any]) -> list[str]:
    order = ["education", "experience", "leadership", "projects", "skills"]
    return [key for key in order if data.get(key) or (key == "skills" and data.get("interests"))]


def _require_mapping(data: dict[str, Any], key: str) -> dict[str, Any]:
    value = data.get(key)
    if not isinstance(value, dict):
        return {}
    return value


def _required_text(item: dict[str, Any], key: str, section: str) -> str:
    value = item.get(key)
    if not value:
        raise CVGeneratorError(f"Missing required field '{key}' in {section}.")
    return str(value).strip()


def _as_list(value: Any) -> list[Any]:
    if value is None:
        return []
    if isinstance(value, list):
        return value
    return [value]


def _prefix(label: str, value: Any) -> str:
    return f"{label} {value}" if value else ""


def _join_sentence_parts(parts: list[Any]) -> str:
    clean = [str(part).strip().rstrip(".") for part in parts if part]
    return ". ".join(clean)


def _stringify_list(value: Any) -> str:
    if isinstance(value, list):
        return ", ".join(str(item).strip() for item in value if str(item).strip())
    return str(value).strip()


def _title_label(value: str) -> str:
    return value.replace("_", " ").title()


def _set_paragraph_layout(paragraph, *, alignment=WD_ALIGN_PARAGRAPH.LEFT, space_before=0, space_after=0) -> None:
    paragraph.alignment = alignment
    fmt = paragraph.paragraph_format
    fmt.left_indent = Inches(0)
    fmt.right_indent = Inches(0)
    fmt.first_line_indent = Inches(0)
    fmt.space_before = Pt(space_before)
    fmt.space_after = Pt(space_after)
    fmt.line_spacing = 1.0
    fmt.line_spacing_rule = WD_LINE_SPACING.SINGLE


def _add_bottom_border(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)

    bottom = p_bdr.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        p_bdr.append(bottom)

    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "000000")
